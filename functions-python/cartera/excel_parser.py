"""Parser del Excel del banco. Origen: admin-app/services/excel_parser.py."""

from __future__ import annotations

from collections import defaultdict
from datetime import datetime
from io import BytesIO
from typing import Any

import openpyxl

from .config import EXCEL_COLUMNS


class ExcelParseError(ValueError):
    """Excel vacío, layout desconocido o sin clientes válidos."""


def safe_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def safe_float(value: Any) -> float:
    if value is None:
        return 0.0
    if isinstance(value, bool):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text:
        return 0.0
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        return float(text)
    except (ValueError, TypeError):
        return 0.0


def safe_int(value: Any) -> int:
    if value is None:
        return 0
    try:
        return int(float(str(value).replace(",", ".")))
    except (ValueError, TypeError):
        return 0


def make_seccion_key(region: str, zona: str, seccion: str) -> str:
    r = str(region).strip() or "SR"
    z = str(zona).strip() or "SZ"
    s = str(seccion).strip().upper() or "SS"
    return f"{r}_{z}_{s}"


def _validate_headers(headers: list[str]) -> None:
    codigo_idx = EXCEL_COLUMNS["codigo_cliente"]
    seccion_idx = EXCEL_COLUMNS["seccion"]
    if codigo_idx >= len(headers) or seccion_idx >= len(headers):
        raise ExcelParseError(
            "El archivo no tiene el layout del banco (faltan columnas)."
        )
    codigo_h = headers[codigo_idx].lower()
    seccion_h = headers[seccion_idx].lower()
    if "codigo" not in codigo_h and "código" not in codigo_h:
        raise ExcelParseError(
            "No se reconoció la columna Código Cliente. "
            f"Título en esa posición: '{headers[codigo_idx]}'."
        )
    if "seccion" not in seccion_h and "sección" not in seccion_h:
        raise ExcelParseError(
            "No se reconoció la columna Sección. "
            f"Título en esa posición: '{headers[seccion_idx]}'."
        )


def _normalize_client(client: dict[str, Any]) -> dict[str, Any]:
    campana = safe_str(client.get("campana"))
    client["campana_banco"] = campana or safe_str(client.get("campana_banco"))
    client["nombre_completo"] = (
        f"{client.get('nombres', '')} "
        f"{client.get('apellido_paterno', '')} "
        f"{client.get('apellido_materno', '')}"
    ).strip()
    seccion = safe_str(client.get("seccion")) or "SIN_SECCION"
    seccion = seccion.upper()
    client["seccion"] = seccion
    client["seccion_key"] = make_seccion_key(
        client.get("region", ""),
        client.get("zona", ""),
        seccion,
    )
    return client


def parse_excel(source: str | bytes | bytearray) -> dict[str, Any]:
    if isinstance(source, (bytes, bytearray)):
        wb = openpyxl.load_workbook(BytesIO(source), data_only=True)
    else:
        wb = openpyxl.load_workbook(source, data_only=True)

    try:
        ws = wb[wb.sheetnames[0]]
        headers = [safe_str(cell.value) for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        _validate_headers(headers)

        all_clients: list[dict[str, Any]] = []
        by_seccion: dict[str, list[dict[str, Any]]] = defaultdict(list)
        codigo_idx = EXCEL_COLUMNS["codigo_cliente"]

        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
            if row is None or codigo_idx >= len(row) or row[codigo_idx] is None:
                continue
            if safe_str(row[codigo_idx]) == "":
                continue

            client: dict[str, Any] = {}
            for field_name, col_index in EXCEL_COLUMNS.items():
                raw_value = row[col_index] if col_index < len(row) else None
                if field_name in ("dias_atraso", "edad"):
                    client[field_name] = safe_int(raw_value)
                elif field_name.startswith("importe_") or field_name.startswith("coordenada_"):
                    client[field_name] = safe_float(raw_value)
                elif field_name.startswith("fecha_"):
                    client[field_name] = safe_str(raw_value)
                else:
                    client[field_name] = safe_str(raw_value)

            _normalize_client(client)
            all_clients.append(client)
            by_seccion[client["seccion_key"]].append(client)

        if not all_clients:
            raise ExcelParseError("El archivo no contiene clientes válidos.")

        total_deuda_asignada = sum(c.get("importe_deuda_asignada", 0) for c in all_clients)
        total_deuda_pendiente = sum(c.get("importe_deuda_pendiente", 0) for c in all_clients)
        summary = {
            "total_clientes": len(all_clients),
            "total_secciones": len(by_seccion),
            "secciones": {k: len(v) for k, v in sorted(by_seccion.items())},
            "total_deuda_asignada": round(total_deuda_asignada, 2),
            "total_deuda_pendiente": round(total_deuda_pendiente, 2),
            "departamentos": sorted(
                {c.get("departamento", "") for c in all_clients if c.get("departamento")}
            ),
        }
        return {
            "all_clients": all_clients,
            "by_seccion": dict(by_seccion),
            "summary": summary,
            "headers": headers,
        }
    finally:
        wb.close()


def get_seccion_summary(by_seccion: dict) -> list[dict[str, Any]]:
    result = []
    for seccion_key in sorted(by_seccion.keys()):
        clients = by_seccion[seccion_key]
        deuda_total = sum(safe_float(c.get("importe_deuda_asignada", 0)) for c in clients)
        deuda_pendiente = sum(safe_float(c.get("importe_deuda_pendiente", 0)) for c in clients)
        departamentos = sorted(
            {c.get("departamento", "") for c in clients if c.get("departamento")}
        )
        letra = seccion_key.rsplit("_", 1)[-1] if seccion_key else seccion_key
        result.append(
            {
                "seccion": seccion_key,
                "seccion_letra": letra,
                "num_clientes": len(clients),
                "deuda_asignada": round(deuda_total, 2),
                "deuda_pendiente": round(deuda_pendiente, 2),
                "departamentos": ", ".join(departamentos),
            }
        )
    return result


def get_hierarchy(all_clients: list) -> dict[str, Any]:
    regions: dict[str, Any] = {}
    for c in all_clients:
        region = safe_str(c.get("region")) or "SIN_REGION"
        zona = safe_str(c.get("zona")) or "SIN_ZONA"
        seccion = safe_str(c.get("seccion")) or "SIN_SECCION"
        deuda_a = safe_float(c.get("importe_deuda_asignada", 0))
        deuda_p = safe_float(c.get("importe_deuda_pendiente", 0))

        if region not in regions:
            regions[region] = {
                "zonas": {},
                "num_clientes": 0,
                "deuda_asignada": 0.0,
                "deuda_pendiente": 0.0,
            }
        r = regions[region]
        r["num_clientes"] += 1
        r["deuda_asignada"] += deuda_a
        r["deuda_pendiente"] += deuda_p

        if zona not in r["zonas"]:
            r["zonas"][zona] = {
                "secciones": {},
                "num_clientes": 0,
                "deuda_asignada": 0.0,
                "deuda_pendiente": 0.0,
            }
        z = r["zonas"][zona]
        z["num_clientes"] += 1
        z["deuda_asignada"] += deuda_a
        z["deuda_pendiente"] += deuda_p

        if seccion not in z["secciones"]:
            z["secciones"][seccion] = {
                "num_clientes": 0,
                "deuda_asignada": 0.0,
                "deuda_pendiente": 0.0,
            }
        s = z["secciones"][seccion]
        s["num_clientes"] += 1
        s["deuda_asignada"] += deuda_a
        s["deuda_pendiente"] += deuda_p

    return {
        "regions": dict(sorted(regions.items())),
        "totals": {
            "num_clientes": sum(r["num_clientes"] for r in regions.values()),
            "deuda_asignada": round(sum(r["deuda_asignada"] for r in regions.values()), 2),
            "deuda_pendiente": round(sum(r["deuda_pendiente"] for r in regions.values()), 2),
        },
    }
